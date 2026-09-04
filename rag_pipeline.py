"""
AI Teacher - RAG Pipeline
Ingests uploaded material (PDF / DOCX / PPTX / TXT) -> chunks -> embeds -> stores in Chroma.
Retrieval returns a plain string, ready to pass as `retrieved_context` into lesson_engine.py.

Requires:
    pip install chromadb sentence-transformers pypdf python-docx python-pptx --break-system-packages

Design notes:
- Uses a local sentence-transformers model (all-MiniLM-L6-v2) for embeddings, so no extra
  API key or cost - good for a hackathon demo where you don't want another dependency
  on rate limits.
- Chroma runs in-process with local persistence (no server to stand up).
- Chunking is simple sentence-window chunking with overlap - good enough for lecture
  notes/textbooks. Swap for a smarter splitter (e.g. header-aware for PPTX) if you have time.
"""

import os
import uuid
import re
from typing import Optional

import chromadb
from chromadb.utils import embedding_functions

CHROMA_DIR = os.environ.get("CHROMA_DIR", "./chroma_store")
CHUNK_SIZE_CHARS = 1200      # ~250-300 tokens, decent context granularity
CHUNK_OVERLAP_CHARS = 200

_embedding_fn = embedding_functions.SentenceTransformerEmbeddingFunction(
    model_name="all-MiniLM-L6-v2"
)
_client = chromadb.PersistentClient(path=CHROMA_DIR)


# ---------------------------------------------------------------------------
# 1. Text extraction per file type
# ---------------------------------------------------------------------------

def extract_text(filepath: str) -> str:
    ext = filepath.lower().rsplit(".", 1)[-1]

    if ext == "pdf":
        from pypdf import PdfReader
        reader = PdfReader(filepath)
        return "\n".join(page.extract_text() or "" for page in reader.pages)

    if ext == "docx":
        import docx
        doc = docx.Document(filepath)
        return "\n".join(p.text for p in doc.paragraphs)

    if ext == "pptx":
        from pptx import Presentation
        prs = Presentation(filepath)
        chunks = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    chunks.append(shape.text_frame.text)
        return "\n".join(chunks)

    if ext == "txt":
        with open(filepath, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()

    raise ValueError(f"Unsupported file type: .{ext}")


# ---------------------------------------------------------------------------
# 2. Chunking - sentence-aware sliding window
# ---------------------------------------------------------------------------

def chunk_text(text: str, chunk_size: int = CHUNK_SIZE_CHARS,
                overlap: int = CHUNK_OVERLAP_CHARS) -> list[str]:
    text = re.sub(r"\s+", " ", text).strip()
    if not text:
        return []

    sentences = re.split(r"(?<=[.!?।])\s+", text)  # includes Hindi danda (।) as a boundary
    chunks = []
    current = ""

    for sentence in sentences:
        if len(current) + len(sentence) <= chunk_size:
            current += (" " if current else "") + sentence
        else:
            if current:
                chunks.append(current)
            # start next chunk with overlap from the end of the previous one
            overlap_text = current[-overlap:] if len(current) > overlap else current
            current = overlap_text + " " + sentence

    if current:
        chunks.append(current)
    return chunks


# ---------------------------------------------------------------------------
# 3. Ingestion - extract, chunk, embed, store
# ---------------------------------------------------------------------------

def ingest_document(filepath: str, doc_id: Optional[str] = None,
                     collection_name: str = "lesson_materials") -> str:
    """
    Ingests a single file into Chroma. Returns the doc_id used, so you can scope
    retrieval to just this document later (useful when multiple students upload
    different files in the same session/collection).
    """
    doc_id = doc_id or str(uuid.uuid4())[:8]
    text = extract_text(filepath)
    chunks = chunk_text(text)

    if not chunks:
        raise ValueError(f"No extractable text found in {filepath}")

    collection = _client.get_or_create_collection(
        name=collection_name, embedding_function=_embedding_fn
    )
    collection.add(
        documents=chunks,
        ids=[f"{doc_id}_{i}" for i in range(len(chunks))],
        metadatas=[{"doc_id": doc_id, "source": os.path.basename(filepath), "chunk_idx": i}
                   for i in range(len(chunks))],
    )
    return doc_id


# ---------------------------------------------------------------------------
# 4. Retrieval - returns a plain string ready for lesson_engine.py
# ---------------------------------------------------------------------------

def retrieve_context(query: str, doc_id: Optional[str] = None, top_k: int = 5,
                      collection_name: str = "lesson_materials") -> str:
    """
    query: what you're teaching right now (e.g. the concept name, or the topic + level).
    doc_id: restrict retrieval to one uploaded document, if the session has one.
    Returns concatenated chunk text, ready to pass as `retrieved_context` to
    plan_lesson() / explain_concept() in lesson_engine.py.
    """
    collection = _client.get_or_create_collection(
        name=collection_name, embedding_function=_embedding_fn
    )
    where = {"doc_id": doc_id} if doc_id else None

    results = collection.query(query_texts=[query], n_results=top_k, where=where)
    docs = results.get("documents", [[]])[0]
    return "\n\n---\n\n".join(docs)


# ---------------------------------------------------------------------------
# Example wiring into lesson_engine.py
# ---------------------------------------------------------------------------
"""
from rag_pipeline import ingest_document, retrieve_context
from lesson_engine import plan_lesson, explain_concept

# On upload:
doc_id = ingest_document("chapter4.pdf")

# When planning:
planning_context = retrieve_context(query=topic, doc_id=doc_id, top_k=8)
state = plan_lesson(topic, level, language, time_budget_min, retrieved_context=planning_context)

# When explaining each concept, re-retrieve scoped to that concept for tighter grounding:
for concept in state.concepts:
    concept_context = retrieve_context(query=concept.name, doc_id=doc_id, top_k=4)
    state.retrieved_context = concept_context   # refresh before explain_concept()
    explanation = explain_concept(state)
"""


if __name__ == "__main__":
    # quick smoke test with a plain text file
    test_path = "sample_notes.txt"
    with open(test_path, "w") as f:
        f.write(
            "Ohm's Law states that current through a conductor is directly proportional "
            "to voltage and inversely proportional to resistance. V = I x R. "
            "If resistance increases while voltage stays constant, current decreases, "
            "not increases - a common student misconception is confusing this with "
            "water flow analogies incorrectly."
        )
    doc_id = ingest_document(test_path)
    print(f"Ingested as doc_id={doc_id}")
    context = retrieve_context("what happens to current when resistance increases", doc_id=doc_id)
    print("\nRetrieved context:\n", context)